"""Exportación del historial de casos para auditoría"""
"""
- Excel: CSV en UTF-8 con BOM, separador ';' y coma decimal (lo abre Excel en español).
- PDF: tabla resumida en A4 apaisado (reportlab).
"""
import csv
import io
from datetime import datetime, timedelta, timezone
import json
from xml.sax.saxutils import escape
from src.registro import COLUMNAS_EXPORTACION
COLUMNAS_PDF = [
    ("fecha_ingreso", "Ingreso", 2.4),
    ("codigo_episodio", "Episodio", 4.4),
    ("dni", "DNI del paciente", 4.2),
    ("motivo_categoria", "Categoría", 2.8),
    ("estado", "Estado", 2.0),
    ("color_final", "Color", 1.9),
    ("decision", "Decisión", 2.0),
    ("motivo_decision", "Motivo de la decisión", 5.0),
    ("profesional_decision", "Decidió", 2.3),
]

COLORES = {
    "ROJO": ("#D32F2F", "#FFFFFF"), "NARANJA": ("#F57C00", "#1A1D26"),
    "AMARILLO": ("#F9C80E", "#1A1D26"), "VERDE": ("#2E7D32", "#FFFFFF"),
    "AZUL": ("#1565C0", "#FFFFFF"),
}

# Presenta las fechas en UTC-03:00, sin depender de la configuración de la PC.
ZONA_INFORME = timezone(timedelta(hours=-3))

TIPOS_EVENTO = {
    "EPISODIO_CREADO": "Inicio de atención",
    "DATOS_INGRESO_REGISTRADOS": "Datos de ingreso revisados",
    "MEDICION_REGISTRADA": "Nueva medición",
    "MEDICION_CORREGIDA": "Corrección de medición",
    "ESTADO_CAMBIADO": "Cambio de estado",
    "DECISION_REGISTRADA": "Decisión profesional",
    "EPISODIO_ANULADO": "Caso anulado (cargado por error)",
}

ESTADOS = {
    "en_ingreso": "En ingreso",
    "en_espera": "En espera",
    "finalizado": "Finalizado",
    "anulado": "Anulado",
}

# Traduce los nombres internos y agrega las unidades de las mediciones.
CAMPOS = {
    "ta_sistolica": ("TA sistólica", "mmHg"),
    "ta_diastolica": ("TA diastólica", "mmHg"),
    "pam": ("Presión arterial media", "mmHg"),
    "frecuencia_cardiaca": ("Frecuencia cardíaca", "lat/min"),
    "frecuencia_respiratoria": ("Frecuencia respiratoria", "resp/min"),
    "saturacion_o2": ("Saturación de oxígeno", "%"),
    "temperatura": ("Temperatura", "°C"),
    "glucemia": ("Glucemia", "mg/dL"),
    "glasgow": ("Glasgow", "/15"),
    "dolor_intensidad": ("Intensidad del dolor", "/10"),
    "estado_conciencia": ("Estado de conciencia", ""),
    "edad": ("Edad", "años"),
    "sexo": ("Sexo", ""),
    "motivo_categoria": ("Categoría del motivo", ""),
    "motivo_consulta": ("Motivo de consulta", ""),
    "dolor_localizacion": ("Localización del dolor", ""),
    "dolor_caracteristicas": ("Características del dolor", ""),
    "sintomas_asociados": ("Síntomas asociados", ""),
    "alergias": ("Alergias", ""),
    "antecedentes_otros": ("Otros antecedentes", ""),
    "tiempo_evolucion_valor": ("Tiempo de evolución", ""),
    "tiempo_evolucion_unidad": ("Unidad del tiempo de evolución", ""),
    "br_sangrado": ("Sangrado", ""),
    "br_convulsion": ("Convulsión", ""),
    "br_disnea": ("Dificultad para respirar", ""),
    "br_trauma": ("Trauma", ""),
    "ant_diabetes": ("Antecedente de diabetes", ""),
    "ant_cardio": ("Antecedente cardiovascular", ""),
    "ant_resp": ("Antecedente respiratorio", ""),
    "ant_anticoag": ("Anticoagulación", ""),
    "ant_inmuno": ("Inmunosupresión", ""),
    "ant_embarazo": ("Embarazo", ""),
    "ant_medicacion": ("Medicación crítica", ""),
}


# Conserva el texto original si la fecha no puede interpretarse.
# No asigna una zona horaria inventada a fechas que no la tienen.
def _fecha_legible(valor):
    if not valor:
        return "No registrada"
    try:
        fecha = datetime.fromisoformat(str(valor).replace("Z", "+00:00"))
        if fecha.utcoffset() is None:
            return f"{valor} (sin zona horaria)"
        return fecha.astimezone(ZONA_INFORME).strftime("%d/%m/%Y %H:%M:%S")
    except ValueError:
        return str(valor)


# Presenta los valores sin True/False ni nombres propios de Python.
def _valor_legible(campo, valor):
    if valor is None or valor == "":
        return "Sin informar"
    if isinstance(valor, bool):
        return "Sí" if valor else "No"

    texto = f"{valor:g}" if isinstance(valor, float) else str(valor)
    unidad = CAMPOS.get(campo, ("", ""))[1]
    return f"{texto} {unidad}".strip()


def _nombre_campo(campo):
    return CAMPOS.get(campo, (campo.replace("_", " ").capitalize(), ""))[0]


# Describe cada evento conservando los cambios y su motivo.
# Los identificadores y versiones siguen guardados en SQLite.
def _detalle_legible(evento):
    tipo = evento.get("tipo")
    detalle = evento.get("detalle") or {}
    valores = detalle.get("valores") or {}
    lineas = []

    if tipo == "EPISODIO_CREADO":
        lineas.append("Se inició la atención.")

    elif tipo in {"DATOS_INGRESO_REGISTRADOS", "MEDICION_REGISTRADA"}:
        for campo, valor in valores.items():
            lineas.append(
                f"{_nombre_campo(campo)}: {_valor_legible(campo, valor)}"
            )

    elif tipo == "MEDICION_CORREGIDA":
        anteriores = detalle.get("valores_anteriores") or {}

        # Incluye también campos retirados o agregados durante la corrección.
        for campo in dict.fromkeys([*anteriores, *valores]):
            antes = anteriores.get(campo)
            despues = valores.get(campo)
            if antes != despues:
                lineas.append(
                    f"{_nombre_campo(campo)}: "
                    f"{_valor_legible(campo, antes)} → "
                    f"{_valor_legible(campo, despues)}"
                )

        referencia = detalle.get("evento_corregido_id")
        if referencia is not None:
            lineas.append(f"Corrige el registro N.º {referencia}.")

    elif tipo == "ESTADO_CAMBIADO":
        antes = detalle.get("estado_anterior")
        despues = detalle.get("estado_nuevo")
        lineas.append(
            f"{ESTADOS.get(antes, antes or 'Sin estado previo')} → "
            f"{ESTADOS.get(despues, despues or 'Sin informar')}"
        )

    elif tipo == "EPISODIO_ANULADO":
        antes = detalle.get("estado_anterior")
        lineas.append(
            f"{ESTADOS.get(antes, antes or 'Sin estado previo')} → Anulado. "
            "El caso no pasó a la lista de espera."
        )

    elif tipo == "DECISION_REGISTRADA":
        decision = detalle.get("decision") or "Sin informar"
        color = detalle.get("color_final") or "Sin color asignado"
        lineas.append(f"Decisión: {decision}. Color final: {color}.")

    else:
        # Evita ocultar información si aparece un tipo de evento nuevo.
        for campo, valor in detalle.items():
            lineas.append(f"{_nombre_campo(campo)}: {valor}")

    if detalle.get("motivo"):
        lineas.append(f"Motivo: {detalle['motivo']}")

    if detalle.get("origen_extraccion"):
        origen = {
            "ia": "IA local, con revisión profesional",
            "simulado": "Caso simulado",
            "manual": "Carga manual",
        }.get(detalle["origen_extraccion"], detalle["origen_extraccion"])
        lineas.append(f"Origen de los datos: {origen}")

    return lineas or ["Sin detalle adicional."]

def _celda(valor, campo):
    if campo == "dni" and (valor is None or valor == ""):
        return "No registrado"
    if valor is None:
        return ""
    if campo.startswith("fecha_") or campo in {"registrado_en", "medido_en"}:
        return _fecha_legible(valor)
    if isinstance(valor, float):
        return f"{valor:g}".replace(".", ",")  # Excel en español usa coma decimal
    texto = str(valor)
    # Evita que Excel interprete el texto como fórmula (inyección CSV).
    if texto[:1] in ("=", "+", "-", "@"):
        texto = "'" + texto
    return texto


# Exporta una fila por evento, repitiendo el resumen del episodio.
# Las columnas de resumen muestran el estado actual.
# Las columnas "Evento:" describen cada registro histórico.
def generar_csv(filas):
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";", lineterminator="\r\n")
        # Incluye el DNI y conserva los códigos para vincular cada atención.
    columnas_resumen = [
        ("dni", "DNI del paciente"),
        *COLUMNAS_EXPORTACION,
    ]
    columnas_evento = [
        ("evento_id", "Evento: ID"),
        ("tipo", "Evento: tipo"),
        ("registrado_en", "Evento: fecha de registro"),
        ("profesional_usuario", "Evento: profesional"),
        ("medido_en", "Evento: fecha de medición"),
        ("evento_corregido_id", "Evento: registro corregido"),
        ("detalle", "Evento: detalle completo"),
    ]

    escritor.writerow(
        [titulo for _, titulo in columnas_resumen]
        + [titulo for _, titulo in columnas_evento]
    )

    for fila in filas:
        resumen = [
            _celda(fila.get(campo), campo)
            for campo, _ in columnas_resumen
        ]

        # Conserva también episodios que no tengan eventos disponibles.
        for evento in fila.get("historial") or [{}]:
            detalle = evento.get("detalle", {})
            datos_evento = {
                "evento_id": evento.get("evento_id"),
                "tipo": TIPOS_EVENTO.get(
                    evento.get("tipo"), evento.get("tipo", "")
                ),
                "registrado_en": evento.get("registrado_en"),
                "profesional_usuario": evento.get("profesional_usuario"),
                "medido_en": detalle.get("medido_en"),
                "detalle": "\n".join(_detalle_legible(evento)),
            }

            escritor.writerow(
                resumen
                + [
                    _celda(datos_evento.get(campo), campo)
                    for campo, _ in columnas_evento
                ]
            )

    # Excel reconoce los acentos mediante el BOM.
    return salida.getvalue().encode("utf-8-sig")


def generar_pdf(filas, filtros_texto, usuario):
    # Import diferido: si falta reportlab, solo falla el PDF (el Excel sigue andando).
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    archivo = io.BytesIO()
    doc = SimpleDocTemplate(
        archivo, pagesize=landscape(A4),
        leftMargin=1 * cm, rightMargin=1 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm,
        title="Historial de casos · Triage HRRG",
    )
    estilos = getSampleStyleSheet()
    chico = estilos["BodyText"].clone("chico", fontSize=7.5, leading=9)
    encabezado = estilos["BodyText"].clone("enc", fontSize=8, leading=10, fontName="Helvetica-Bold")

        # Identifica explícitamente la zona horaria del informe.
    ahora = datetime.now(ZONA_INFORME).strftime("%d/%m/%Y %H:%M") + " (UTC-03:00)"
    partes = [
        Paragraph("Historial de casos · Triage asistido HRRG", estilos["Title"]),
        Paragraph(
            f"Exportado por <b>{usuario}</b> el {ahora} · {filtros_texto} · {len(filas)} episodios · "
                        "<b>DNI enmascarado</b> (solo los últimos 3 números) · Archivo protegido con la contraseña de quien exportó · Datos ficticios del prototipo.",
            chico,
        ),
        Spacer(1, 0.3 * cm),
    ]

    datos = [[Paragraph(t, encabezado) for _, t, _ in COLUMNAS_PDF]]
    estilo_tabla = [
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0B4BE")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EAF0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    col_color = [c for c, _, _ in COLUMNAS_PDF].index("color_final")
    for i, fila in enumerate(filas, start=1):
        datos.append([
            Paragraph(_celda(fila.get(c), c).lstrip("'").replace("&", "&amp;").replace("<", "&lt;"), chico)
            for c, _, _ in COLUMNAS_PDF
        ])
        fondo, tinta = COLORES.get(fila.get("color_final") or "", (None, None))
        if fondo:
            estilo_tabla.append(("BACKGROUND", (col_color, i), (col_color, i), colors.HexColor(fondo)))
            estilo_tabla.append(("TEXTCOLOR", (col_color, i), (col_color, i), colors.HexColor(tinta)))

    tabla = Table(datos, colWidths=[a * cm for _, _, a in COLUMNAS_PDF], repeatRows=1)
    tabla.setStyle(TableStyle(estilo_tabla))
    partes.append(tabla)
        # Presenta la evolución con títulos claros y fechas uniformes.
    for fila in filas:
        partes.append(Spacer(1, 0.5 * cm))
        partes.append(Paragraph(
            "Historial de la atención "
            + escape(str(fila.get("codigo_episodio", ""))),
            estilos["Heading2"],
        ))

                # Identifica al paciente también al comienzo de su historial.
        partes.append(Paragraph(
            "DNI del paciente: " + escape(_celda(fila.get("dni"), "dni")),
            chico,
        ))

        historial = fila.get("historial") or []
        if not historial:
            partes.append(Paragraph("Sin eventos disponibles.", chico))
            continue

        for evento in historial:
            detalle = evento.get("detalle") or {}
            tipo = TIPOS_EVENTO.get(
                evento.get("tipo"), evento.get("tipo", "Evento")
            )

            # Mantiene una referencia breve para vincular las correcciones.
            titulo = f"{tipo} · Registro N.º {evento.get('evento_id', '')}"
            partes.append(Paragraph(escape(titulo), estilos["Heading3"]))

            registro = _fecha_legible(evento.get("registrado_en"))
            profesional = evento.get("profesional_usuario") or "No registrado"
            partes.append(Paragraph(
                escape(f"Registrado: {registro} · Profesional: {profesional}"),
                chico,
            ))

            # Distingue cuándo se tomó el dato de cuándo se registró o corrigió.
            if detalle.get("medido_en"):
                partes.append(Paragraph(
                    escape("Fecha de la toma: "
                           + _fecha_legible(detalle["medido_en"])),
                    chico,
                ))

            for linea in _detalle_legible(evento):
                partes.append(Paragraph(escape(linea), chico))

            partes.append(Spacer(1, 0.2 * cm))
    def pie(canvas, documento):
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(landscape(A4)[0] - 1 * cm, 0.6 * cm, f"Página {documento.page}")

    doc.build(partes, onFirstPage=pie, onLaterPages=pie)
    return archivo.getvalue()


# ---------------------------------------------------------------------------
# [Propuesta Barby 26/10 · a revisar por Maricel]
# Excel de verdad (.xlsx) con tres hojas:
#   "Resumen" → una fila por episodio (para filtrar, contar y hacer tablas dinámicas).
#   "Eventos" → una fila por evento (el detalle completo, con fecha y profesional).
#   "Info"    → quién exportó, cuándo, con qué filtros y avisos.
# Números y fechas quedan como números y fechas de Excel (no como texto).
# La columna DNI aparece solo si la ruta la agregó a las filas (así, si el grupo decide
# exportar sin DNI, alcanza con no agregarla en app.py).
# ---------------------------------------------------------------------------
def _fecha_excel(valor):
    """ISO con zona → datetime sin zona, en UTC-03:00 (Excel no guarda zonas)."""
    if not valor:
        return None
    try:
        fecha = datetime.fromisoformat(str(valor).replace("Z", "+00:00"))
    except ValueError:
        return str(valor)
    if fecha.utcoffset() is None:
        return fecha
    return fecha.astimezone(ZONA_INFORME).replace(tzinfo=None)


def generar_xlsx(filas, filtros_texto, usuario):
    # Import diferido: si falta openpyxl, solo falla el Excel (el PDF sigue andando).
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    con_dni = any("dni" in fila for fila in filas)
    libro = Workbook()
    negrita = Font(bold=True)
    fondo_titulo = PatternFill("solid", fgColor="E8EAF0")
    arriba = Alignment(vertical="top", wrap_text=True)

    def escribir(hoja, fila, columna, valor, formato=None):
        celda = hoja.cell(row=fila, column=columna, value=valor)
        # Un texto que empieza con "=" no debe ejecutarse como fórmula en Excel.
        if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@"):
            celda.data_type = "s"
        if formato:
            celda.number_format = formato
        return celda

    def encabezado(hoja, titulos, anchos):
        for col, titulo in enumerate(titulos, start=1):
            c = hoja.cell(row=1, column=col, value=titulo)
            c.font, c.fill = negrita, fondo_titulo
            c.alignment = Alignment(vertical="center", wrap_text=True)
            hoja.column_dimensions[get_column_letter(col)].width = anchos[col - 1]
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}1"

    # ---------- Hoja 1: Resumen (una fila por episodio) ----------
    resumen = libro.active
    resumen.title = "Resumen"
    columnas = ([("dni", "DNI del paciente")] if con_dni else []) + list(COLUMNAS_EXPORTACION)
    anchos = []
    for campo, titulo in columnas:
        if campo.startswith("fecha_"):
            anchos.append(17)
        elif campo in ("motivo_consulta", "motivo_decision", "antecedentes", "banderas_rojas"):
            anchos.append(34)
        elif campo in ("codigo_episodio", "codigo_paciente"):
            anchos.append(26)
        else:
            anchos.append(max(10, min(20, len(titulo) + 2)))
    encabezado(resumen, [t for _, t in columnas], anchos)
    resumen.row_dimensions[1].height = 30

    col_color = [c for c, _ in columnas].index("color_final") + 1
    for i, fila in enumerate(filas, start=2):
        for j, (campo, _) in enumerate(columnas, start=1):
            valor = fila.get(campo)
            if campo.startswith("fecha_"):
                escribir(resumen, i, j, _fecha_excel(valor), "dd/mm/yyyy hh:mm")
            elif campo == "dni":
                escribir(resumen, i, j, valor or "No registrado")
            else:
                escribir(resumen, i, j, valor)
        fondo, tinta = COLORES.get(fila.get("color_final") or "", (None, None))
        if fondo:
            c = resumen.cell(row=i, column=col_color)
            c.fill = PatternFill("solid", fgColor=fondo.lstrip("#"))
            c.font = Font(bold=True, color=tinta.lstrip("#"))

    # ---------- Hoja 2: Eventos (una fila por evento) ----------
    eventos = libro.create_sheet("Eventos")
    titulos_ev = (
        ["Episodio"] + (["DNI del paciente"] if con_dni else []) + ["Código de paciente",
         "Registro N.º", "Tipo de evento", "Fecha de registro", "Profesional",
         "Fecha de la toma", "Corrige el registro N.º", "Detalle"]
    )
    anchos_ev = [26] + ([16] if con_dni else []) + [26, 11, 24, 17, 16, 17, 12, 70]
    encabezado(eventos, titulos_ev, anchos_ev)
    n = 2
    for fila in filas:
        for evento in fila.get("historial") or []:
            detalle = evento.get("detalle") or {}
            valores = (
                [fila.get("codigo_episodio")]
                + ([fila.get("dni") or "No registrado"] if con_dni else [])
                + [
                    fila.get("codigo_paciente"),
                    evento.get("evento_id"),
                    TIPOS_EVENTO.get(evento.get("tipo"), evento.get("tipo")),
                    _fecha_excel(evento.get("registrado_en")),
                    evento.get("profesional_usuario"),
                    _fecha_excel(detalle.get("medido_en")),
                    detalle.get("evento_corregido_id"),
                    "\n".join(_detalle_legible(evento)),
                ]
            )
            for j, valor in enumerate(valores, start=1):
                formato = "dd/mm/yyyy hh:mm:ss" if hasattr(valor, "year") else None
                escribir(eventos, n, j, valor, formato).alignment = arriba
            n += 1

    # ---------- Hoja 3: Info ----------
    info = libro.create_sheet("Info")
    ahora = datetime.now(ZONA_INFORME).strftime("%d/%m/%Y %H:%M") + " (UTC-03:00)"
    renglones = [
        ("Historial de casos · Triage asistido HRRG", None),
        ("Exportado por", usuario),
        ("Fecha de exportación", ahora),
        ("Filtros", filtros_texto),
        ("Episodios", len(filas)),
        ("Eventos", n - 2),
        ("Horas", "Todas las fechas están en UTC-03:00 (hora de Argentina)."),
        ("Identificación", "DNI enmascarado: solo se ven los últimos 3 números. El archivo se abre con la contraseña de quien exportó."
                           if con_dni else "Sin DNI: el paciente figura con su código interno."),
        ("Datos", "Datos ficticios del prototipo (no son pacientes reales)."),
    ]
    for i, (clave, valor) in enumerate(renglones, start=1):
        c = info.cell(row=i, column=1, value=clave)
        c.font = Font(bold=True, size=14 if i == 1 else 11)
        if valor is not None:
            escribir(info, i, 2, valor)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 80

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# [Propuesta Barby 26/10 · a revisar por Maricel]
# Protección del archivo exportado:
#   - El DNI sale enmascarado: solo se ven los últimos 3 números (••.•••.222).
#   - El archivo se cifra con la contraseña del auditor que exporta
#     (Excel: cifrado nativo de Office con msoffcrypto; PDF: AES-256 con pikepdf).
#     Sin esa contraseña no se puede abrir. No se guarda en ningún lado.
# ---------------------------------------------------------------------------
def enmascarar_dni(dni):
    if not dni:
        return None
    digitos = "".join(c for c in str(dni) if c.isdigit())
    if len(digitos) < 4:
        return "•" * len(digitos)
    oculto = "•" * (len(digitos) - 3) + digitos[-3:]
    # Mismo formato de puntos que el DNI: 8 dígitos → ••.•••.222 · 7 dígitos → •.•••.222
    partes = []
    while oculto:
        partes.insert(0, oculto[-3:])
        oculto = oculto[:-3]
    return ".".join(partes)


def proteger_xlsx(contenido, clave):
    # Import diferido: si falta msoffcrypto-tool, solo falla la protección del Excel.
    from msoffcrypto.format.ooxml import OOXMLFile
    entrada, salida = io.BytesIO(contenido), io.BytesIO()
    archivo = OOXMLFile(entrada)
    archivo.encrypt(clave, salida)
    return salida.getvalue()


def proteger_pdf(contenido, clave):
    import pikepdf  # import diferido: si falta, solo falla la protección del PDF
    salida = io.BytesIO()
    with pikepdf.open(io.BytesIO(contenido)) as pdf:
        pdf.save(salida, encryption=pikepdf.Encryption(user=clave, owner=clave, R=6))
    return salida.getvalue()
